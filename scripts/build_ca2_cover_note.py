"""Build CA02 cover note with the same cover-page layout as CA01."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deliverables" / "B9AI105_Group_CA02_20082900_20093736_10592825_20098301_Cover_Note.docx"

NAVY = RGBColor(0x00, 0x00, 0x00)
SLATE = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x33, 0x33, 0x33)


def _run(run, *, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _p(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    color=None,
    center=False,
    space_before=0,
    space_after=6,
    line_spacing=1.15,
):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = line_spacing
    r = p.add_run(text)
    _run(r, size=size, bold=bold, italic=italic, color=color or SLATE)
    return p


def _heading(doc, text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True


def _bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        _run(r, bold=True, color=NAVY)
        r = p.add_run(text)
        _run(r, color=SLATE)
    else:
        r = p.add_run(text)
        _run(r, color=SLATE)


def _add_ca01_style_cover(doc: Document) -> None:
    """Page 1 — same fields / layout pattern as the CA01 title page."""
    # Top spacer so the block sits like a cover page
    for _ in range(6):
        _p(doc, "", size=12, space_after=6)

    _p(doc, "Dublin Business School", size=16, bold=True, center=True, space_after=10)
    _p(doc, "B9AI105 Reinforcement Learning", size=14, bold=True, center=True, space_after=8)
    _p(
        doc,
        "Master of Science in Artificial Intelligence - Group A",
        size=12,
        center=True,
        space_after=8,
    )
    _p(doc, "Lecturer Name: Dr. Oleksandr Bezrukavyi", size=12, center=True, space_after=18)

    _p(doc, "CA02 – Group Assessment", size=16, bold=True, center=True, space_after=6)
    _p(
        doc,
        "GréineQ — Reinforcement Learning for Solar Battery Dispatch",
        size=12,
        italic=True,
        center=True,
        space_after=18,
    )

    _p(doc, "Group Members:", size=12, bold=True, center=True, space_after=6)
    _p(doc, "Nathan Lucio - 20082900", size=12, center=True, space_after=2)
    _p(doc, "Nadeesha Jayasuriya - 20093736", size=12, center=True, space_after=2)
    _p(doc, "Emmanuel Addoh - 10592825", size=12, center=True, space_after=2)
    _p(doc, "Bahadir Demir - 20098301", size=12, center=True, space_after=18)

    _p(doc, "Submission Date: 11th August 2026", size=12, center=True, space_after=6)

    # Page break before extension body
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    _add_ca01_style_cover(doc)

    # ----- Body (extension note) -----
    _heading(doc, "1. Repository, deployed system, and presentation")
    _bullet(doc, "https://github.com/NathanLCR/rl-battery-dispatch", bold_prefix="GitHub: ")
    _bullet(doc, "https://greineq-agent.sudocod.com/", bold_prefix="Deployed Web Twin: ")
    _bullet(
        doc,
        "deliverables/presentation/GreineQ_CA2_Presentation.pptx",
        bold_prefix="Final presentation (PPTX): ",
    )
    _bullet(
        doc,
        "https://drive.google.com/drive/folders/12vc-RRwxKix4hWDKO1u6VhppgT2KjyNz?usp=sharing",
        bold_prefix="Presentation + recording (Google Drive): ",
    )
    _bullet(doc, "deliverables/GreineQ_User_Guide.md", bold_prefix="User guide: ")
    _p(
        doc,
        "Local run: from rl-battery-dispatch/ use .\\run_web.ps1 (http://localhost:8080) or Docker "
        "(see Dockerfile/README.md). Trained models belong in results/models/.",
        size=10,
        italic=True,
        color=MUTED,
        space_after=10,
    )

    _heading(doc, "2. Relationship to CA01")
    _p(
        doc,
        "This CA02 submission is a direct extension of our CA01 group project (GréineQ). "
        "CA01 built the core tabular RL microgrid: Ausgrid solar + AEMO NSW1 prices, a custom MDP, "
        "NumPy Q-Learning / SARSA, greedy and rule baselines, and a Streamlit replay dashboard.",
        space_after=8,
    )
    _p(doc, "CA01 in brief", bold=True, space_after=4)
    _bullet(doc, "Solar-only battery control with three actions: hold, solar-charge, discharge.")
    _bullet(doc, "Discretised state space (324 states in the CA01 formulation).")
    _bullet(doc, "Monetised reward based on grid-import cost (and related battery-aware terms).")
    _bullet(
        doc,
        "Finding: RL beat a weak tertile rule (~33% lower import cost) but did not beat greedy "
        "self-consumption on held-out test days — strong baselines matter.",
    )
    _bullet(doc, "Deliverable: Streamlit dashboard for day replay and controller comparison.")

    _heading(doc, "3. What CA02 added (extensions)")
    _p(
        doc,
        "CA02 extends the same codebase and research question into an arbitrage-capable, "
        "forecast-information experiment and a deployable digital twin.",
        space_after=8,
    )
    _bullet(
        doc,
        "expanded from 3 to 5 actions (hold, solar-charge, discharge, grid-charge, export).",
        bold_prefix="Action space: ",
    )
    _bullet(
        doc,
        "experimental wholesale-exposed export tariff (not a normal household FiT), "
        "with efficiencies, cycling cost, and terminal SOC valuation.",
        bold_prefix="Tariff / physics: ",
    )
    _bullet(
        doc,
        "does a coarse three-bin true 4-hour price-direction feature help tabular Q-Learning "
        "beat a strong greedy 5-action baseline?",
        bold_prefix="Research question: ",
    )
    _bullet(
        doc,
        "Current Q (540 states) vs Privileged Q (1620 states) vs Greedy vs no-battery vs "
        "perfect-foresight bound; 53 held-out days × 5 seeds.",
        bold_prefix="Controllers / evaluation: ",
    )
    _bullet(
        doc,
        "Greedy remained best deployable (~AUD 90.39 vs no-battery 160.75). "
        "Privileged true-direction foresight did not help on average (Current Q ~124.23; Privileged ~139.68).",
        bold_prefix="Headline result: ",
    )
    _bullet(
        doc,
        "FastAPI Web Twin (published): Overview, Digital Twin, Price Monitor, Agent Play, Results; "
        "plus Streamlit research twin.",
        bold_prefix="Product / demo: ",
    )
    _bullet(
        doc,
        "human vs RL vs Greedy on identical day twins (oversight demo with disclosed information asymmetry).",
        bold_prefix="Play vs Agent: ",
    )
    _bullet(
        doc,
        "Double Q-Learning support, forecast modules, live AEMO price monitor, Docker packaging, "
        "user guide and presentation pack.",
        bold_prefix="Engineering extras: ",
    )

    _heading(doc, "4. How to read this submission")
    _bullet(doc, "Page 1 is the formal cover sheet (same layout pattern as CA01).")
    _bullet(doc, "This note links CA01 → CA02 and lists repository / demo / presentation links.")
    _bullet(doc, "Full experiment write-up: deliverables/notebooklm/CA2_Forecast_Info_Experiment_Findings.md")
    _bullet(doc, "Code entry points: src/ (MDP + agents + experiment), webapp/ (deployed UI), dashboard/ (Streamlit)")
    _bullet(doc, "Reproduce headline experiment: see README.md (oracle / forecast_exp commands)")
    _p(
        doc,
        "GréineQ remains a software digital twin for research and demonstration — not a physical battery controller.",
        size=10,
        italic=True,
        color=MUTED,
        space_after=6,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
