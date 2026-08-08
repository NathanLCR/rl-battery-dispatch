"""Build a polished GréineQ User Guide Word document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
AMBER = RGBColor(0xD9, 0x77, 0x06)


def _set_run(run, *, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _docx_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = SLATE
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)


def _add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("GréineQ")
    _set_run(run, size=28, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    run = sub.add_run("User Guide")
    _set_run(run, size=20, bold=True, color=AMBER)

    tag = doc.add_paragraph()
    run = tag.add_run(
        "Operate and explore the residential battery digital twin: compare controllers "
        "on historical days, play against a trained agent, check live wholesale prices, "
        "and review the CA2 experiment headline."
    )
    _set_run(run, size=11, color=SLATE)

    meta = doc.add_paragraph()
    run = meta.add_run("Published app: ")
    _set_run(run, size=10, bold=True, color=MUTED)
    run = meta.add_run("https://greineq-agent.sudocod.com/")
    _set_run(run, size=10, color=SLATE)

    brand = doc.add_paragraph()
    run = brand.add_run("Brand spelling: GréineQ")
    _set_run(run, size=10, italic=True, color=MUTED)
    run = brand.add_run("  ·  Filenames may use ASCII GreineQ")
    _set_run(run, size=10, italic=True, color=MUTED)

    note = doc.add_paragraph()
    run = note.add_run(
        "Production may lag the latest ca2_agent branch — run locally for the newest Twin / Play / Results UI."
    )
    _set_run(run, size=10, italic=True, color=MUTED)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY if level == 1 else SLATE
        run.font.name = "Calibri"


def _para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run(run, size=11, bold=bold, italic=italic, color=SLATE)


def _bullet(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        _set_run(run, size=11, bold=True, color=NAVY)
        run = p.add_run(text)
        _set_run(run, size=11, color=SLATE)
    else:
        run = p.add_run(text)
        _set_run(run, size=11, color=SLATE)


def _numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    _set_run(run, size=11, color=SLATE)


def _code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = NAVY


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_run(run, size=10, bold=True, color=NAVY)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(val)
            _set_run(run, size=10, color=SLATE)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def build_user_guide_docx() -> Path:
    doc = Document()
    _docx_styles(doc)
    _add_title_block(doc)

    # --- 1 ---
    _heading(doc, "1. Open the dashboard")
    _heading(doc, "Online", level=2)
    _para(
        doc,
        "Open the published URL on the cover page. For demos, prefer a local run so you control models and UI version.",
    )

    _heading(doc, "Local (recommended for demos)", level=2)
    _para(doc, "From the repository root rl-battery-dispatch/:")
    _code_block(
        doc,
        [
            "pip install -r requirements.txt",
            ".\\run_dashboard.ps1",
        ],
    )
    _para(doc, "Or:")
    _code_block(
        doc,
        [
            "python -m streamlit run dashboard/app.py --server.port 8501",
        ],
    )
    _para(doc, "Then open http://localhost:8501")
    _bullet(doc, "Smoke check (optional): python scripts/system_test.py")

    # --- 2 ---
    _heading(doc, "2. Navigation")
    _para(doc, "Use the top bar to switch views:")
    _add_table(
        doc,
        ["View", "What it does"],
        [
            ["Digital Twin", "Same-day replay of several controllers; KPIs, charts, timestep inspector"],
            ["Price Monitor", "Live AEMO NSW1 price + typical PV/load (informational)"],
            ["Agent Play", "You vs RL vs Greedy for 48 half-hour steps"],
            ["Results", "Curated CA2 experiment chart and interpretation"],
            ["Overview", "Landing page and heuristic animation preview"],
        ],
        col_widths=[1.6, 4.8],
    )
    _para(
        doc,
        "The Overview animation is a heuristic preview, not the trained Q-Learning agent.",
        italic=True,
    )

    # --- 3 ---
    _heading(doc, "3. Digital Twin")
    _para(
        doc,
        "Historical simulation only — one calendar day, 48 half-hour intervals. "
        "Controllers see the same PV, load, and prices for that day.",
    )

    _heading(doc, "Sidebar — Simulation setup", level=2)
    _numbered(doc, "Day split — Test (default), Validation, or Train.")
    _numbered(doc, "Simulation day — pick a date from that split. Demo suggestion: 2012-07-14.")

    _heading(doc, "Sidebar — Controllers", level=2)
    _para(doc, "Toggle which policies to compare (defaults are usually on):")
    _add_table(
        doc,
        ["Controller", "Meaning"],
        [
            [
                "Greedy (5-action)",
                "Self-use surplus solar; price-timed grid-charge / export",
            ],
            [
                "Current Q",
                "Tabular Q-Learning using current state + price only",
            ],
            [
                "Privileged Q — true 4h signal",
                "Q-Learning plus the true next-four-hour three-bin price direction "
                "(information probe, not a realistic forecast)",
            ],
        ],
        col_widths=[2.3, 4.1],
    )

    _heading(doc, "Run comparison", level=2)
    _bullet(doc, "First open auto-runs once so the page is not blank.")
    _bullet(doc, "After you change day, split, or controllers, click Run comparison again.")
    _bullet(doc, "If settings changed but you have not re-run, a banner asks you to run again.")

    _heading(doc, "Experiment settings (collapsed)", level=2)
    _para(doc, "Optional extras:")
    _bullet(doc, "Reward function: battery-aware (default) vs cost-only")
    _bullet(doc, "Extra baselines: solar-only greedy, tertile rule")
    _bullet(doc, "Additional RL: SARSA, Double Q-Learning")
    _bullet(doc, "Model pickers when those agents are enabled")
    _bullet(doc, "Q-value diagnostics in the download section")
    _para(
        doc,
        "Trained weights live under results/models/. If none are present, RL checkboxes will error until you train.",
    )

    _heading(doc, "Main results layout", level=2)
    _numbered(doc, "Context bar — mode, day, preview policy")
    _numbered(doc, "Winning agent strip + four KPIs (net cost, savings vs no-battery, gap to best RL, final SOC)")
    _numbered(doc, "Why did X win? — short explanation + perfect-foresight cost bound")
    _numbered(doc, "Controller comparison — net cost, savings %, final SOC bars (lower cost is better)")
    _numbered(doc, "Detailed metrics — Cost / Energy / Battery tabs")
    _numbered(doc, "Charts — cost, energy, and battery trajectories")
    _numbered(doc, "Timestep inspector — step through 1–48; action, SOC, and step cost per controller")
    _numbered(doc, "Diagnostics & download — step log; optional Q-values; JSON download")
    _para(doc, "Live prices are not on this page — use Price Monitor.", bold=True)

    # --- 4 ---
    _heading(doc, "4. Live Price Monitor")
    _para(doc, "Separate from the historical Twin.")
    _add_table(
        doc,
        ["Field", "Source"],
        [
            ["Wholesale / retail estimate", "Live AEMO NSW1"],
            [
                "Typical PV / load",
                "Historical medians for the current hour — not a live household meter",
            ],
        ],
        col_widths=[2.2, 4.2],
    )
    _bullet(doc, "If the feed is down, you will see a notice; Digital Twin still works on historical data.")
    _bullet(
        doc,
        "An optional Current Q recommendation uses live price with illustrative SOC and typical PV/load — "
        "for illustration only, not a metered home.",
    )

    # --- 5 ---
    _heading(doc, "5. Agent Play (Play vs Agent)")
    _para(
        doc,
        "Interactive oversight demo: you operate one battery; Current Q (or Privileged Q) and Greedy "
        "run identical twins in parallel. Lowest adjusted electricity cost wins.",
    )

    _heading(doc, "Game setup", level=2)
    _numbered(doc, "Choose a test day.")
    _numbered(doc, "Choose opponent: Current Q or Privileged Q — true 4h signal.")
    _numbered(doc, "Choose a trained model.")
    _numbered(doc, "Click Start game.")

    callout = doc.add_paragraph()
    run = callout.add_run("Information asymmetry (important): ")
    _set_run(run, size=11, bold=True, color=AMBER)
    run = callout.add_run(
        "You always see a realistic 4-hour price forecast. Privileged Q may use a true four-hour "
        "direction feature. This is an oversight / teaching demo, not a scientifically fair equal-information contest."
    )
    _set_run(run, size=11, color=SLATE)

    _heading(doc, "During the game", level=2)
    _bullet(doc, "Four state cards: price, SOC, solar, load.")
    _bullet(doc, "Choose your action (or keys 1–5):")
    _add_table(
        doc,
        ["#", "Action", "Intent"],
        [
            ["1", "Hold", "No battery movement"],
            ["2", "Solar charge", "Store surplus solar"],
            ["3", "Discharge", "Supply household demand"],
            ["4", "Grid charge", "Buy from the grid for later"],
            ["5", "Export", "Sell stored electricity"],
        ],
        col_widths=[0.6, 1.6, 4.2],
    )
    _bullet(doc, "Unavailable actions show a disable reason (e.g. empty battery, no surplus solar).")
    _bullet(doc, "Scoreboard: You / RL / Greedy running costs.")
    _bullet(doc, "Tabs: forecast view and agent decision insight.")
    _bullet(doc, "Pause / Quit game (or sidebar Quit to setup).")

    _heading(doc, "End of day", level=2)
    _para(
        doc,
        "Summary of who won, optional replay of decisions, try another day, and download of the match data.",
    )

    # --- 6 ---
    _heading(doc, "6. Experiment Results")
    _para(
        doc,
        "Read-only CA2 headline for the wholesale-export, true-direction privileged probe:",
    )
    _bullet(doc, "KPI strip and bar chart (lower net cost is better)")
    _bullet(doc, "Compact results table (mean ± std for Q agents across five seeds)")
    _bullet(doc, "Collapsed Experiment setup and optional per-seed CSV")

    takeaway = doc.add_paragraph()
    run = takeaway.add_run("Takeaway under this setup: ")
    _set_run(run, size=11, bold=True, color=NAVY)
    run = takeaway.add_run(
        "Greedy remained the best evaluated deployable controller; the three-bin true direction "
        "feature did not improve Q-Learning on average."
    )
    _set_run(run, size=11, color=SLATE)
    _para(
        doc,
        "Do not mix these totals with older fixed-FiT experiment numbers.",
        italic=True,
    )

    # --- 7 ---
    _heading(doc, "7. Controllers & metrics (quick glossary)")
    _heading(doc, "Controllers", level=2)
    _add_table(
        doc,
        ["Name", "Role"],
        [
            ["No-battery reference", "Cost if there were no battery"],
            ["Perfect foresight / oracle", "Information lower-cost bound"],
            ["Greedy (5-action)", "Strong rule baseline (often best deployable)"],
            ["Current Q", "Learns from current discretised state only"],
            ["Privileged Q", "Same + true 4h three-bin direction (headline probe)"],
        ],
        col_widths=[2.3, 4.1],
    )

    _heading(doc, "Key metrics", level=2)
    _add_table(
        doc,
        ["Metric", "Meaning"],
        [
            [
                "Net cost (AUD)",
                "Imports − export revenue + cycling / terminal effects (lower is better)",
            ],
            ["Savings", "Improvement vs no-battery reference"],
            ["Export revenue", "Income under the configured export tariff"],
            ["Grid-charge cost", "Cost of deliberate grid charging"],
            ["Final SOC", "End-of-day battery level (condition, not the win criterion)"],
            ["Throughput", "How much energy moved through the battery"],
        ],
        col_widths=[1.8, 4.6],
    )
    _para(
        doc,
        "Reward in training = negative adjusted cost (agents maximise reward by minimising cost).",
    )

    _heading(doc, "Tariff note", level=2)
    _para(
        doc,
        "Export uses an experimental wholesale-exposed price (AEMO AUD/MWh ÷ 1000), not a typical "
        "Australian household FiT. Round-trip sell is hard; valuable foresight is mostly "
        "“buy cheap → serve load later.”",
    )

    # --- 8 ---
    _heading(doc, "8. Suggested demo path (~10 minutes)")
    _numbered(doc, "Overview — brand and what the twin is.")
    _numbered(doc, "Results — headline chart and “greedy still wins.”")
    _numbered(
        doc,
        "Digital Twin — day 2012-07-14, default controllers, Run comparison; open winner explanation; "
        "step the inspector through a peak-price window.",
    )
    _numbered(
        doc,
        "Agent Play — short stretch of decisions vs Current Q (call out forecast vs privileged if you switch opponents).",
    )
    _numbered(doc, "Price Monitor — live vs historical twin distinction (optional).")
    _para(
        doc,
        "Keep verified models under results/models/ for the demo; avoid mid-demo retrain.",
        bold=True,
    )

    # --- 9 ---
    _heading(doc, "9. Troubleshooting")
    _add_table(
        doc,
        ["Issue", "What to do"],
        [
            ["Blank Twin / stale numbers", "Click Run comparison after changing settings"],
            ["No trained models", "Train or copy .npy models into results/models/"],
            ["Incomplete day warning", "Choose another day with full 48 intervals"],
            ["Live feed unavailable", "Expected on some networks; Twin still works"],
            ["Charts missing", "Ensure plotly is installed (requirements.txt)"],
            ["Port in use", "Change --server.port or stop the other Streamlit process"],
        ],
        col_widths=[2.4, 4.0],
    )

    # --- 10 ---
    _heading(doc, "10. Scope & ethics")
    _bullet(doc, "Software twin and research demo — not a physical battery controller.")
    _bullet(doc, "Live Price Monitor mixes live wholesale with typical PV/load.")
    _bullet(doc, "Play vs Privileged Q is an oversight framing, not a fair RL contest.")
    _bullet(doc, "Wholesale export tariff is experimental and disclosed in the write-up.")
    _para(
        doc,
        "For experiment design and findings, see deliverables/notebooklm/CA2_Forecast_Info_Experiment_Findings.md "
        "and the presentation pack under deliverables/presentation/.",
        italic=True,
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("GréineQ  ·  B9AI105  ·  Dashboard User Guide")
    _set_run(run, size=9, italic=True, color=MUTED)

    path = OUT / "GreineQ_User_Guide.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    out = build_user_guide_docx()
    print(f"Wrote {out}")
